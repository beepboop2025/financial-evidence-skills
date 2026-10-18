"""Compare native office output with every expected cell, including inert text."""
import xml.etree.ElementTree as ET
from zipfile import ZipFile

NS = {
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
}


def attr(prefix, name):
    return "{" + NS[prefix] + "}" + name


def paragraph(node):
    text = node.text or ""
    for child in node:
        if child.tag == attr("text", "s"):
            text += " " * int(child.get(attr("text", "c"), "1"))
        elif child.tag == attr("text", "tab"):
            text += "\t"
        elif child.tag == attr("text", "line-break"):
            text += "\n"
        else:
            text += paragraph(child)
        text += child.tail or ""
    return text


def verify(path, expected):
    """Return a bounded all-cell acceptance count; raise on any mismatch."""
    checked = 0
    if path.suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, data_only=False)
        assert book.sheetnames == [s["name"] for s in expected]
        for sheet in expected:
            actual = book[sheet["name"]]
            for r, values in enumerate(sheet["rows"], 1):
                for c, wanted in enumerate(values, 1):
                    cell = actual.cell(r, c)
                    assert cell.data_type != "f", (sheet["name"], r, c, "unexpected formula")
                    observed = "" if cell.value is None else cell.value
                    assert observed == wanted and (type(observed) is bool) == (type(wanted) is bool), (sheet["name"], r, c, observed, wanted)
                    checked += 1
        assert not book._external_links, "Unexpected external workbook links"
        return checked
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        root = ET.fromstring(archive.read("content.xml"))
    tables = root.findall(".//table:table", NS)
    assert [t.get(attr("table", "name")) for t in tables] == [s["name"] for s in expected]
    for table, sheet in zip(tables, expected):
        rows = []
        for row in table.findall("table:table-row", NS):
            rows.extend([row] * min(int(row.get(attr("table", "number-rows-repeated"), "1")), len(sheet["rows"])))
            if len(rows) >= len(sheet["rows"]):
                break
        assert len(rows) >= len(sheet["rows"])
        for r, wanted_row in enumerate(sheet["rows"]):
            observed = []
            for cell in rows[r].findall("table:table-cell", NS):
                if len(observed) >= len(wanted_row):
                    break
                wanted = wanted_row[len(observed)]
                formula = cell.get(attr("table", "formula"))
                if formula is not None:
                    assert type(wanted) is bool and formula == ("of:=TRUE()" if wanted else "of:=FALSE()"), "Unexpected formula"
                kind = cell.get(attr("office", "value-type"))
                if kind == "float":
                    value = float(cell.get(attr("office", "value")))
                elif kind == "boolean":
                    raw = cell.get(attr("office", "boolean-value"))
                    assert raw in ("true", "false", "1", "0"), ("Invalid XML boolean", raw)
                    value = raw in ("true", "1")
                else:
                    value = "\n".join(paragraph(p) for p in cell.findall("text:p", NS))
                count = min(int(cell.get(attr("table", "number-columns-repeated"), "1")), len(wanted_row) - len(observed))
                observed.extend([value] * count)
            observed += [""] * (len(wanted_row) - len(observed))
            assert observed == wanted_row, (sheet["name"], r, observed, wanted_row)
            assert [type(v) is bool for v in observed] == [type(v) is bool for v in wanted_row]
            checked += len(wanted_row)
    return checked
